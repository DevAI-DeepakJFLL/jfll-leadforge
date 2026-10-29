import os
import re
import csv
import json
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict

# Paths
ACTIVE_EXCEL_PATH = "Backend/data/Active Customer List For the Last 3 Month.xlsx"
MASTER_BACKUP_EXCEL_PATH = "Backend/data/.backups/2026-10-17-pre-customer-dedup/Customer with contact.xlsx"
OLD_ACTIVE_CSV_PATH = "Backend/data/processed/active_customers_leads.csv"
DATAMYNE_PARQUET_PATH = "Backend/data/processed/datamyne_leads_zstd.parquet"

OUTPUT_ACTIVE_PARQUET = "Backend/data/processed/active_customers_leads_zstd.parquet"
OUTPUT_ACTIVE_CSV = "Backend/data/processed/active_customers_leads.csv"
OUTPUT_INACTIVE_PARQUET = "Backend/data/processed/inactive_customers_leads_zstd.parquet"
OUTPUT_INACTIVE_CSV = "Backend/data/processed/inactive_customers_leads.csv"
OUTPUT_RAW_EXCEL = "Backend/data/raw/Customer with contact.xlsx"
OUTPUT_AUDIT_CSV = "Backend/data/raw/DUPLICATES_REMOVED_AUDIT_Customer_with_contact.csv"

FRONTEND_ACTIVE_TS = "Frontend/src/data/jfllCustomersData.ts"
FRONTEND_INACTIVE_TS = "Frontend/src/data/jfllInactiveCustomersData.ts"


def normalize_company(s):
    if not s or pd.isna(s):
        return ""
    s = str(s).upper().strip()
    s = re.sub(r"^(M/S\.?|MESSRS\.?|M/S\s+|MS\.)\s*", "", s)
    s = re.sub(r"&", " AND ", s)
    s = re.sub(r"\bPVT\.?\s*LTD\.?\b", "PRIVATE LIMITED", s)
    s = re.sub(r"\bP\.?\s*LTD\.?\b", "PRIVATE LIMITED", s)
    s = re.sub(r"\bPVT\.?\b", "PRIVATE", s)
    s = re.sub(r"\bLTD\.?\b", "LIMITED", s)
    s = re.sub(r"\bCO\.?\b", "COMPANY", s)
    s = re.sub(r"\bCORP\.?\b", "CORPORATION", s)
    s = re.sub(r"\bINC\.?\b", "INCORPORATED", s)
    s = re.sub(r"\(PROP\..*?\)", " ", s)
    s = re.sub(r"\(PROP\s+.*?\)", " ", s)
    s = re.sub(r"\(INDIA\)", "INDIA", s)
    s = re.sub(r"\(I\)", "I", s)
    s = re.sub(r"\(HUF\)", "HUF", s)
    s = re.sub(r"[^A-Z0-9]", "", s)
    return s


def clean_corp_stem(s):
    if not s:
        return ""
    s = normalize_company(s)
    changed = True
    while changed:
        old_s = s
        for suffix in [
            "PRIVATELIMITED",
            "PRIVATE",
            "LIMITED",
            "LLP",
            "INCORPORATED",
            "CORPORATION",
            "COMPANY",
            "HUF",
            "INDIA",
            "PVTLTD",
        ]:
            if s.endswith(suffix):
                s = s[:-len(suffix)]
        changed = (s != old_s)
    return s


def extract_phones(p):
    if not p or pd.isna(p):
        return set()
    raw = str(p)
    clean = re.sub(r"[^\d/,\- ]", "", raw)
    tokens = re.split(r"[/,;\n\r|]", clean)
    phones = set()
    for tok in tokens:
        digits = re.sub(r"\D", "", tok)
        if len(digits) == 10:
            if not all(d == digits[0] for d in digits) and digits != "1234567890":
                phones.add(digits)
        elif len(digits) > 10:
            if digits.startswith("91") and len(digits) == 12:
                phones.add(digits[2:])
            elif digits.startswith("0") and len(digits) == 11:
                phones.add(digits[1:])
            else:
                phones.add(digits[-10:])
    return phones


def extract_emails(e):
    if not e or pd.isna(e):
        return set()
    raw = str(e).lower().strip()
    found = set(re.findall(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}", raw))
    return {
        x
        for x in found
        if not any(
            d in x for d in ["test.com", "example.com", "none@", "na@", "noemail@"]
        )
    }


def format_phone_display(raw):
    if not raw or pd.isna(raw):
        return None
    raw_str = str(raw).strip()
    if not raw_str or raw_str.lower() in ["none", "nan", "null"]:
        return None
    if raw_str.startswith("+"):
        return raw_str
    # Extract 10-digit number
    phones = extract_phones(raw_str)
    if phones:
        first = sorted(list(phones))[0]
        return f"+91-{first}"
    return raw_str


def main():
    print("=" * 80)
    print("REVAMPING JFLL CUSTOMER DATASET (3-MONTH INACTIVITY THRESHOLD)")
    print("=" * 80)

    # 1. Load Datasets
    print(f"\n[Step 1/6] Loading source datasets...")
    df_new_active = pd.read_excel(ACTIVE_EXCEL_PATH)
    print(f"  • New Active Customers (Last 3 Months): {len(df_new_active)} records")

    df_master = pd.read_excel(MASTER_BACKUP_EXCEL_PATH)
    print(f"  • Master Customer Database (Pre-dedup Backup): {len(df_master)} records")

    df_old_active = pd.read_csv(OLD_ACTIVE_CSV_PATH)
    print(f"  • Previous Active Customers CSV: {len(df_old_active)} records")

    df_dm = pd.read_parquet(DATAMYNE_PARQUET_PATH)
    print(f"  • DataMyne Leads: {len(df_dm)} records")

    # Index Reference Datasets for Fast Lookup
    print("\n[Step 2/6] Building lookup indexes...")
    old_by_norm = {}
    old_by_stem = {}
    for _, r in df_old_active.iterrows():
        n = normalize_company(r["Full Name"])
        if n and n not in old_by_norm:
            old_by_norm[n] = r
        st = clean_corp_stem(r["Full Name"])
        if len(st) >= 4 and st not in old_by_stem:
            old_by_stem[st] = r

    master_by_norm = {}
    master_by_stem = {}
    for _, r in df_master.iterrows():
        n = normalize_company(r["Full Name"])
        if n and n not in master_by_norm:
            master_by_norm[n] = r
        st = clean_corp_stem(r["Full Name"])
        if len(st) >= 4 and st not in master_by_stem:
            master_by_stem[st] = r

    dm_by_norm = {}
    dm_by_stem = {}
    dm_by_phone = defaultdict(list)
    dm_by_email = defaultdict(list)
    for _, r in df_dm.iterrows():
        n = normalize_company(r["company_name"])
        if n and n not in dm_by_norm:
            dm_by_norm[n] = r
        st = clean_corp_stem(r["company_name"])
        if len(st) >= 4 and st not in dm_by_stem:
            dm_by_stem[st] = r
        for p in extract_phones(r.get("phone", "")):
            dm_by_phone[p].append(r["company_name"])
        for em in extract_emails(r.get("email", "")):
            dm_by_email[em].append(r["company_name"])

    # 3. Process Active Customers (578 records)
    print("\n[Step 3/6] Processing 578 Active Customer accounts...")
    active_records = []
    for idx, r in df_new_active.iterrows():
        name = str(r["Full Name"]).strip()
        credit_status = str(r["Credit Status"]).strip() if pd.notna(r["Credit Status"]) else "approved"
        sales_person = str(r["Sales Person"]).strip() if pd.notna(r["Sales Person"]) and str(r["Sales Person"]).strip() != "" else None
        business_last_3m = round(float(r["Total Business Received From Cusomer In Last 3 Month"]), 2) if pd.notna(r["Total Business Received From Cusomer In Last 3 Month"]) else 0.0

        n = normalize_company(name)
        st = clean_corp_stem(name)

        # Contact and metadata matching priority: Old Active -> Master -> DataMyne
        ref = None
        if n in old_by_norm:
            ref = old_by_norm[n]
        elif len(st) >= 4 and st in old_by_stem:
            ref = old_by_stem[st]
        elif n in master_by_norm:
            ref = master_by_norm[n]
        elif len(st) >= 4 and st in master_by_stem:
            ref = master_by_stem[st]

        dm_ref = None
        if ref is None:
            if n in dm_by_norm:
                dm_ref = dm_by_norm[n]
            elif len(st) >= 4 and st in dm_by_stem:
                dm_ref = dm_by_stem[st]

        # Extract contact person
        contact_person = None
        if ref is not None:
            if "Contact Person Name" in ref and pd.notna(ref["Contact Person Name"]) and str(ref["Contact Person Name"]).strip():
                contact_person = str(ref["Contact Person Name"]).strip()
            elif "Contact Person Full Name" in ref and pd.notna(ref["Contact Person Full Name"]) and str(ref["Contact Person Full Name"]).strip():
                contact_person = str(ref["Contact Person Full Name"]).strip()

        # Extract mobile number
        mobile = None
        if ref is not None and "Mobile Number" in ref and pd.notna(ref["Mobile Number"]):
            mobile = format_phone_display(ref["Mobile Number"])
        elif dm_ref is not None and pd.notna(dm_ref.get("phone")):
            mobile = format_phone_display(dm_ref["phone"])

        # Extract email
        email = None
        if ref is not None and "Email" in ref and pd.notna(ref["Email"]):
            raw_em = str(ref["Email"]).strip()
            if raw_em and raw_em.lower() not in ["none", "nan", "null"]:
                email = raw_em
        elif dm_ref is not None and pd.notna(dm_ref.get("email")):
            raw_em = str(dm_ref["email"]).strip()
            if raw_em and raw_em.lower() not in ["none", "nan", "null"]:
                # take first email if multiple
                email = [x.strip() for x in raw_em.split(",") if "@" in x][0] if "@" in raw_em else None

        # Extract Customer Group
        customer_group = None
        if ref is not None and "Customer Group" in ref and pd.notna(ref["Customer Group"]) and str(ref["Customer Group"]).strip():
            customer_group = str(ref["Customer Group"]).strip()
        else:
            upper = name.upper()
            if any(w in upper for w in ["LOGISTICS", "FREIGHT", "EXPRESS", "CARGO", "FORWARD", "SHIPPING", "LINES", "CLEARING", "AIRWAYS", "TRAVEL", "TRANSPORT", "AVIA"]):
                customer_group = "Subagent"
            elif any(w in upper for w in ["GLOBAL", "OVERSEAS", "INTERNATIONAL", "WORLDWIDE", "TRANS", "SDN BHD", "S.R.L.", "GMBH", "B.V.", "LLC"]):
                customer_group = "Overseas Agent"
            else:
                customer_group = "Shipper / Consignee"

        # Extract Entity Type
        entity_type = None
        if ref is not None and "Type" in ref and pd.notna(ref["Type"]) and str(ref["Type"]).strip():
            entity_type = str(ref["Type"]).strip()
        else:
            upper = name.upper()
            if "PVT LTD" in upper or "PRIVATE LIMITED" in upper or "PVT. LTD" in upper:
                entity_type = "Private Limited"
            elif "LLP" in upper:
                entity_type = "LLP"
            elif "LTD" in upper or "LIMITED" in upper:
                entity_type = "Public Limited"
            elif any(w in upper for w in ["ENTERPRISES", "TRADING", "TRADERS", "AGENCY"]):
                entity_type = "Individual / Proprietorship"
            elif "PARTNERSHIP" in upper:
                entity_type = "Partnership"
            else:
                entity_type = "Company"

        active_records.append({
            "id": f"jfll-{idx + 1}",
            "Full Name": name,
            "Territory": "All Territories",
            "Contact Person Name": contact_person,
            "Mobile Number": mobile,
            "Email": email,
            "Status": "Active",
            "Customer Group": customer_group,
            "Type": entity_type,
            "Credit Status": credit_status,
            "Sales Person": sales_person,
            "Total Business Received From Cusomer In Last 3 Month": business_last_3m,
        })

    df_active_processed = pd.DataFrame(active_records)
    print(f"  • Processed Active Customers: {len(df_active_processed)} rows")
    print(f"    - With contact person: {df_active_processed['Contact Person Name'].notna().sum()}")
    print(f"    - With phone: {df_active_processed['Mobile Number'].notna().sum()}")
    print(f"    - With email: {df_active_processed['Email'].notna().sum()}")
    print(f"    - Customer Groups breakdown:\n{df_active_processed['Customer Group'].value_counts().to_string()}")

    # 4. Partition Master Records into Active Purge and Inactive Customers
    print("\n[Step 4/6] Partitioning 2,778 Master Records into Active (Purged) and Inactive (>3m inactivity)...")
    # Index new active names
    active_name_set = set(df_active_processed["Full Name"].apply(normalize_company))
    active_stem_set = {clean_corp_stem(n) for n in df_active_processed["Full Name"] if len(clean_corp_stem(n)) >= 4}

    audit_purged_active = []
    kept_inactive_records = []

    for idx, r in df_master.iterrows():
        r_name = str(r["Full Name"]).strip() if pd.notna(r["Full Name"]) else ""
        r_contact = str(r["Contact Person Full Name"]).strip() if pd.notna(r["Contact Person Full Name"]) else None
        r_phone = str(r["Mobile Number"]).strip() if pd.notna(r["Mobile Number"]) else None
        r_email = str(r["Email"]).strip() if pd.notna(r["Email"]) else None

        r_norm = normalize_company(r_name)
        r_stem = clean_corp_stem(r_name)
        r_phones = extract_phones(r_phone)
        r_emails = extract_emails(r_email)

        # Check if matches active
        is_active = (r_norm in active_name_set) or (len(r_stem) >= 4 and r_stem in active_stem_set)

        if is_active:
            audit_purged_active.append({
                "Original_Row_Number": idx + 2,
                "Company_Name": r_name,
                "Contact_Person": r_contact or "",
                "Mobile_Number": r_phone or "",
                "Email": r_email or "",
                "Purge_Category": "Active JFLL Customer (Last 3 Months)",
                "Purge_Reasons": f"Active Customer Match ({r_name})",
            })
        else:
            # Check DataMyne cross-match
            dm_reasons = []
            if r_norm in dm_by_norm:
                dm_reasons.append(f"Name Match ({dm_by_norm[r_norm]['company_name']})")
            elif len(r_stem) >= 4 and r_stem in dm_by_stem:
                dm_reasons.append(f"Entity Stem Match ({dm_by_stem[r_stem]['company_name']})")

            for p in r_phones:
                if p in dm_by_phone:
                    dm_reasons.append(f"Phone Match ({p} -> {dm_by_phone[p][0]})")
                    break

            for em in r_emails:
                if em in dm_by_email:
                    dm_reasons.append(f"Email Match ({em} -> {dm_by_email[em][0]})")
                    break

            is_in_dm = len(dm_reasons) > 0
            kept_inactive_records.append({
                "Full Name": r_name,
                "Contact Person Full Name": r_contact,
                "Mobile Number": format_phone_display(r_phone) if r_phone else None,
                "Email": r_email,
                "in_datamyne": is_in_dm,
                "dm_match_details": " | ".join(dm_reasons) if is_in_dm else "",
            })

    print(f"  • Active accounts purged from prospecting: {len(audit_purged_active)}")
    print(f"  • Inactive accounts retained (>3 months inactivity): {len(kept_inactive_records)}")
    in_dm_count = sum(1 for r in kept_inactive_records if r["in_datamyne"])
    print(f"    - Matched in DataMyne (Verified Shippers): {in_dm_count}")
    print(f"    - Archive / Dormant Records: {len(kept_inactive_records) - in_dm_count}")

    # Build Inactive DataFrame
    inactive_records = []
    for idx, r in enumerate(kept_inactive_records):
        name = r["Full Name"]
        upper = name.upper()
        if "PVT LTD" in upper or "PRIVATE LIMITED" in upper or "PVT. LTD" in upper:
            ent_type = "Private Limited"
        elif "LLP" in upper:
            ent_type = "LLP"
        elif "LTD" in upper or "LIMITED" in upper:
            ent_type = "Public Limited"
        elif any(w in upper for w in ["ENTERPRISES", "TRADING", "TRADERS", "AGENCY"]):
            ent_type = "Individual / Proprietorship"
        elif "PARTNERSHIP" in upper:
            ent_type = "Partnership"
        else:
            ent_type = "Company"

        inactive_records.append({
            "id": f"jfll-inact-{idx + 1:04d}",
            "Full Name": name,
            "Territory": "All Territories",
            "Contact Person Name": r["Contact Person Full Name"],
            "Mobile Number": r["Mobile Number"],
            "Email": r["Email"],
            "Status": "Not Active",
            "Customer Group": "Inactive (DataMyne Match)" if r["in_datamyne"] else "Inactive Customer",
            "Type": ent_type,
            "in_datamyne": r["in_datamyne"],
            "datamyne_match_details": r["dm_match_details"],
        })

    df_inactive_processed = pd.DataFrame(inactive_records)

    # 5. Export Processed Parquet, CSV, Audit Trail, and Excel per xlsx skill
    print("\n[Step 5/6] Exporting Parquet, CSV, and formatted Excel files...")

    # Active Parquet & CSV (columns matching backend model schema + new intelligence fields)
    active_export_cols = [
        "Full Name",
        "Territory",
        "Contact Person Name",
        "Mobile Number",
        "Email",
        "Status",
        "Customer Group",
        "Type",
        "Credit Status",
        "Sales Person",
        "Total Business Received From Cusomer In Last 3 Month",
    ]
    df_active_processed[active_export_cols].to_parquet(
        OUTPUT_ACTIVE_PARQUET, engine="pyarrow", compression="zstd", index=False
    )
    df_active_processed[active_export_cols].to_csv(
        OUTPUT_ACTIVE_CSV, index=False, encoding="utf-8-sig"
    )
    print(f"  ✓ Saved Active Customers Parquet: {OUTPUT_ACTIVE_PARQUET} ({len(df_active_processed)} rows)")
    print(f"  ✓ Saved Active Customers CSV: {OUTPUT_ACTIVE_CSV} ({len(df_active_processed)} rows)")

    # Inactive Parquet & CSV
    df_inactive_processed.to_parquet(
        OUTPUT_INACTIVE_PARQUET, engine="pyarrow", compression="zstd", index=False
    )
    df_inactive_processed.to_csv(
        OUTPUT_INACTIVE_CSV, index=False, encoding="utf-8-sig"
    )
    print(f"  ✓ Saved Inactive Customers Parquet: {OUTPUT_INACTIVE_PARQUET} ({len(df_inactive_processed)} rows)")
    print(f"  ✓ Saved Inactive Customers CSV: {OUTPUT_INACTIVE_CSV} ({len(df_inactive_processed)} rows)")

    # Audit CSV for purged active records
    with open(OUTPUT_AUDIT_CSV, "w", newline="", encoding="utf-8-sig") as f:
        fieldnames = [
            "Original_Row_Number",
            "Company_Name",
            "Contact_Person",
            "Mobile_Number",
            "Email",
            "Purge_Category",
            "Purge_Reasons",
        ]
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(audit_purged_active)
    print(f"  ✓ Saved Purge Audit Log: {OUTPUT_AUDIT_CSV} ({len(audit_purged_active)} rows)")

    # Format Backend/data/raw/Customer with contact.xlsx with 2,279 rows per xlsx skill guidelines
    print(f"  • Formatting Excel file {OUTPUT_RAW_EXCEL} with openpyxl (Calibri 10pt, professional styling)...")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Customer with contact"
    ws.views.sheetView[0].showGridLines = True

    # Styling constants
    FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    FILL_HEADER = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")  # Professional Navy
    FONT_DATA = Font(name="Calibri", size=10, bold=False, color="000000")
    ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
    THIN_BORDER_SIDE = Side(border_style="thin", color="D9D9D9")
    CELL_BORDER = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE)

    headers = ["Full Name", "Contact Person Full Name", "Mobile Number", "Email"]
    ws.append(headers)

    # Style header row
    ws.row_dimensions[1].height = 26
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Populate rows
    for r_idx, rec in enumerate(kept_inactive_records, start=2):
        ws.row_dimensions[r_idx].height = 20
        c1 = ws.cell(row=r_idx, column=1, value=rec["Full Name"])
        c2 = ws.cell(row=r_idx, column=2, value=rec["Contact Person Full Name"] or "")
        c3 = ws.cell(row=r_idx, column=3, value=rec["Mobile Number"] or "")
        c4 = ws.cell(row=r_idx, column=4, value=rec["Email"] or "")

        for cell in (c1, c2, c3, c4):
            cell.font = FONT_DATA
            cell.alignment = ALIGN_LEFT
            cell.border = CELL_BORDER

    # Set deliberate column widths (auto-fit with safety margin)
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 16)

    # Freeze header row
    ws.freeze_panes = "A2"

    wb.save(OUTPUT_RAW_EXCEL)
    print(f"  ✓ Saved formatted Excel file: {OUTPUT_RAW_EXCEL} ({len(kept_inactive_records)} data rows)")

    # 6. Generate TypeScript Files for Frontend
    print("\n[Step 6/6] Generating TypeScript frontend data files...")

    # Active Customers TS
    active_ts_items = []
    for r in active_records:
        active_ts_items.append({
            "id": r["id"],
            "fullName": r["Full Name"],
            "territory": r["Territory"],
            "contactPersonName": r["Contact Person Name"],
            "mobileNumber": r["Mobile Number"],
            "email": r["Email"],
            "status": "Active",
            "customerGroup": r["Customer Group"],
            "type": r["Type"],
            "saved": False,
            "creditStatus": r["Credit Status"],
            "salesPerson": r["Sales Person"],
            "businessLast3m": r["Total Business Received From Cusomer In Last 3 Month"],
        })

    active_ts_content = f"""// {len(active_ts_items)} Active Customer Records from Jet Freight Logistics Master Database (Last 3 Months Activity Window)
import {{ JfllCustomer }} from '@/types/jfllCustomer';

export const JFLL_CUSTOMERS_DATA: JfllCustomer[] = {json.dumps(active_ts_items, indent=2)};
"""
    with open(FRONTEND_ACTIVE_TS, "w", encoding="utf-8") as f:
        f.write(active_ts_content)
    print(f"  ✓ Saved Frontend Active TS: {FRONTEND_ACTIVE_TS} ({len(active_ts_items)} records)")

    # Inactive Customers TS
    inactive_ts_items = []
    for r in inactive_records:
        inactive_ts_items.append({
            "id": r["id"],
            "fullName": r["Full Name"],
            "territory": r["Territory"],
            "contactPersonName": r["Contact Person Name"],
            "mobileNumber": r["Mobile Number"],
            "email": r["Email"],
            "status": "Not Active",
            "customerGroup": r["Customer Group"],
            "type": r["Type"],
            "saved": False,
            "inDataMyne": r["in_datamyne"],
        })

    inactive_ts_content = f"""// {len(inactive_ts_items):,} Inactive Customer Records from Customer with contact.xlsx (>3 Months Inactivity Window)
import {{ JfllCustomer }} from '@/types/jfllCustomer';

export const JFLL_INACTIVE_CUSTOMERS_DATA: JfllCustomer[] = {json.dumps(inactive_ts_items, indent=2)};
"""
    with open(FRONTEND_INACTIVE_TS, "w", encoding="utf-8") as f:
        f.write(inactive_ts_content)
    print(f"  ✓ Saved Frontend Inactive TS: {FRONTEND_INACTIVE_TS} ({len(inactive_ts_items)} records)")

    print("\n" + "=" * 80)
    print("ETL REVAMP COMPLETED SUCCESSFULLY!")
    print(f"• Total Active Customers (3-Month Activity): {len(active_ts_items)}")
    print(f"• Total Inactive Customers (>3-Month Inactivity): {len(inactive_ts_items)}")
    print(f"• Total Master Directory Accounts: {len(active_ts_items) + len(inactive_ts_items)}")
    print("=" * 80)


if __name__ == "__main__":
    main()

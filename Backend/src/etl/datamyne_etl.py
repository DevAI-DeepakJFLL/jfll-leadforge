"""
ETL Pipeline for DataMyne Leads Dataset (Valid & Missing Contacts).
Aggregates, standardizes, enriches, and exports DataMyne leads into ZSTD-compressed Parquet.
Produces:
- Backend/data/processed/datamyne_valid_leads_zstd.parquet (16,899 records)
- Backend/data/processed/datamyne_missing_leads_zstd.parquet (17,791 records)
- Backend/data/processed/datamyne_leads_zstd.parquet (34,690 unified records)
- Backend/data/processed/datamyne_leads.csv
"""

import time
from pathlib import Path

import openpyxl
import pandas as pd


def run_datamyne_etl():
    start_time = time.time()
    print("=" * 70)
    print("STARTING DATAMYNE PARQUET (.ZSTD) ETL PIPELINE")
    print("=" * 70)

    base_dir = Path(__file__).resolve().parent.parent.parent
    datamyne_dir = base_dir / "data" / "Datamyne"
    output_dir = base_dir / "data" / "processed"
    output_dir.mkdir(parents=True, exist_ok=True)

    valid_file = datamyne_dir / "Valid Contacts" / "Master_Cleaned_Contacts_Directory.xlsx"
    missing_file = datamyne_dir / "Missing Contacts" / "Master_Cleaned_Missing_Contacts_Directory.xlsx"

    print(f"Reading Valid Contacts from: {valid_file.name}...")
    wb_val = openpyxl.load_workbook(valid_file, read_only=True)
    ws_val = wb_val["All Cleaned Contacts"]
    val_rows = list(ws_val.iter_rows(values_only=True))
    val_header = val_rows[0]
    val_data = val_rows[1:]
    df_val = pd.DataFrame(val_data, columns=val_header)
    df_val["dataset_type"] = "valid"
    print(f"Loaded {len(df_val)} valid contact records.")

    print(f"Reading Missing Contacts from: {missing_file.name}...")
    wb_mis = openpyxl.load_workbook(missing_file, read_only=True)
    ws_mis = wb_mis["All Cleaned Missing Contacts"]
    mis_rows = list(ws_mis.iter_rows(values_only=True))
    mis_header = mis_rows[0]
    mis_data = mis_rows[1:]
    df_mis = pd.DataFrame(mis_data, columns=mis_header)
    df_mis["dataset_type"] = "missing"
    print(f"Loaded {len(df_mis)} missing contact records.")

    # Drop 'Name Variations Merged' if present
    for df in [df_val, df_mis]:
        if "Name Variations Merged" in df.columns:
            df.drop(columns=["Name Variations Merged"], inplace=True)

    # Standardize column names
    col_mapping = {
        "Company Name": "company_name",
        "Role": "role",
        "Transport Mode": "transport_mode",
        "Category": "category",
        "Subcategory": "subcategory",
        "Contact Type": "contact_type",
        "Phone": "phone",
        "Email": "email",
        "Tax ID (IEC/PAN)": "tax_id",
        "City": "city",
        "Address": "address",
        "PIN Code": "pin_code",
        "Directors / Contact": "directors",
        "Trade Lanes / Sectors": "trade_lanes",
        "Top HS Codes": "top_hs_codes",
        "Total Shipments": "total_shipments",
        "Total FOB/CIF (USD)": "total_value_usd",
        "Latest Activity Date": "latest_activity_date"
    }

    df_val.rename(columns=col_mapping, inplace=True)
    df_mis.rename(columns=col_mapping, inplace=True)

    # Generate sequential unique IDs
    df_val["id"] = [f"dm-val-{i+1:05d}" for i in range(len(df_val))]
    df_mis["id"] = [f"dm-mis-{i+1:05d}" for i in range(len(df_mis))]

    # Combine into unified dataset
    unified_df = pd.concat([df_val, df_mis], ignore_index=True)

    # Enforce data types & cleanup
    for df in [df_val, df_mis, unified_df]:
        df["country"] = "India"
        df["total_shipments"] = pd.to_numeric(df["total_shipments"], errors="coerce").fillna(0).astype("int64")
        df["total_value_usd"] = pd.to_numeric(df["total_value_usd"], errors="coerce").fillna(0.0).astype("float64")
        df["company_name"] = df["company_name"].astype(str).str.strip()
        df["role"] = df["role"].astype(str).str.strip()
        df["transport_mode"] = df["transport_mode"].astype(str).str.strip().replace(["None", "nan", ""], "Air")
        df["category"] = df["category"].astype(str).str.strip()
        df["subcategory"] = df["subcategory"].astype(str).str.strip()
        df["contact_type"] = df["contact_type"].astype(str).str.strip()
        
        # String fields handle nulls
        for str_col in ["phone", "email", "tax_id", "city", "address", "pin_code", "directors", "trade_lanes", "top_hs_codes", "latest_activity_date"]:
            df[str_col] = df[str_col].apply(lambda x: str(x).strip() if pd.notna(x) and str(x).strip() not in ["None", "nan", ""] else "")

    target_columns = [
        "id",
        "company_name",
        "country",
        "role",
        "transport_mode",
        "category",
        "subcategory",
        "contact_type",
        "phone",
        "email",
        "tax_id",
        "city",
        "address",
        "pin_code",
        "directors",
        "trade_lanes",
        "top_hs_codes",
        "total_shipments",
        "total_value_usd",
        "latest_activity_date",
        "dataset_type"
    ]

    df_val = df_val[target_columns]
    df_mis = df_mis[target_columns]
    unified_df = unified_df[target_columns]

    # Save to Parquet with ZSTD compression
    val_parquet_path = output_dir / "datamyne_valid_leads_zstd.parquet"
    mis_parquet_path = output_dir / "datamyne_missing_leads_zstd.parquet"
    unified_parquet_path = output_dir / "datamyne_leads_zstd.parquet"
    unified_csv_path = output_dir / "datamyne_leads.csv"

    print("\nWriting ZSTD compressed Parquet files...")
    df_val.to_parquet(val_parquet_path, engine="pyarrow", compression="zstd", index=False)
    print(f"  ✓ {val_parquet_path.name}: {len(df_val):,} rows ({val_parquet_path.stat().st_size / 1024:.1f} KB)")

    df_mis.to_parquet(mis_parquet_path, engine="pyarrow", compression="zstd", index=False)
    print(f"  ✓ {mis_parquet_path.name}: {len(df_mis):,} rows ({mis_parquet_path.stat().st_size / 1024:.1f} KB)")

    unified_df.to_parquet(unified_parquet_path, engine="pyarrow", compression="zstd", index=False)
    print(f"  ✓ {unified_parquet_path.name}: {len(unified_df):,} rows ({unified_parquet_path.stat().st_size / 1024:.1f} KB)")

    unified_df.to_csv(unified_csv_path, index=False, encoding="utf-8-sig")
    print(f"  ✓ {unified_csv_path.name}: {len(unified_df):,} rows ({unified_csv_path.stat().st_size / (1024*1024):.2f} MB)")

    print("\n" + "=" * 70)
    print(f"DATAMYNE ETL COMPLETED IN {time.time()-start_time:.2f}s")
    print("=" * 70)
    return unified_df

if __name__ == "__main__":
    run_datamyne_etl()

import os
import re
import glob
import time
from collections import Counter, defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DATAMYNE_DIR = "Backend/data/Datamyne"
VALID_MASTER_PATH = os.path.join(DATAMYNE_DIR, "Valid Contacts/Master_Cleaned_Contacts_Directory.xlsx")
MISSING_MASTER_PATH = os.path.join(DATAMYNE_DIR, "Missing Contacts/Master_Cleaned_Missing_Contacts_Directory.xlsx")

# 2-Digit Chapter to Logistics Vertical & Subcategory Map
CHAPTER_MAP = {
    # Perishables
    "02": ("Perishables", "Meat & Edible Meat Offal"),
    "03": ("Perishables", "Seafood & Marine Products"),
    "04": ("Perishables", "Dairy & Edible Animal Products"),
    "06": ("Perishables", "Floriculture & Cut Flowers"),
    "07": ("Perishables", "Fresh Vegetables & Produce"),
    "08": ("Perishables", "Fresh Fruits"),
    "09": ("Perishables", "Spices, Coffee & Tea"),
    "12": ("Perishables", "Agricultural Seeds & Sowing"),
    "13": ("Perishables", "Botanical Extracts & Resins"),
    "20": ("Perishables", "Processed Fruits & Vegetables"),
    "21": ("Perishables", "Processed Foods & Edible Preparations"),
    "22": ("Perishables", "Beverages & Liquids"),

    # Pharma & Healthcare
    "30": ("Pharma & Healthcare", "Pharmaceutical Formulations"),
    "90": ("Pharma & Healthcare", "Medical & Surgical Devices"),

    # Chemicals & APIs
    "28": ("Chemicals & Materials", "Inorganic Chemicals"),
    "29": ("Chemicals & Materials", "Organic Chemicals & APIs"),
    "38": ("Chemicals & Materials", "Specialty Chemicals & Auxiliaries"),

    # Textiles & Apparel
    "50": ("Textiles & Apparel", "Silk & Silk Textiles"),
    "51": ("Textiles & Apparel", "Wool & Animal Hair"),
    "52": ("Textiles & Apparel", "Cotton & Cotton Textiles"),
    "54": ("Textiles & Apparel", "Synthetic Filaments & Yarns"),
    "55": ("Textiles & Apparel", "Man-made Staple Fibres"),
    "56": ("Textiles & Apparel", "Wadding, Felt & Non-wovens"),
    "58": ("Textiles & Apparel", "Special Woven Fabrics"),
    "59": ("Textiles & Apparel", "Fabrics & Technical Textiles"),
    "60": ("Textiles & Apparel", "Knitted or Crocheted Fabrics"),
    "61": ("Textiles & Apparel", "Apparel & Ready-Made Garments (Knitted)"),
    "62": ("Textiles & Apparel", "Apparel & Ready-Made Garments (Woven)"),
    "63": ("Textiles & Apparel", "Made-up Textile Articles"),

    # Engineering & Machinery
    "84": ("Engineering & Machinery", "Industrial Machinery & Boilers"),
    "73": ("Engineering & Machinery", "Iron, Steel & Hardware"),
    "82": ("Engineering & Machinery", "Tools & Implements"),

    # Electronics & Electrical
    "85": ("Electronics & Electrical", "Electrical Equipment & Switchgear"),

    # Automotive & EV
    "87": ("Automotive & EV", "Motor Vehicles, EV & Auto Parts"),

    # General Cargo
    "39": ("General Cargo", "Plastics & Polymers"),
    "40": ("General Cargo", "Rubber & Rubber Products"),
    "48": ("General Cargo", "Paper, Paperboard & Packaging"),
    "70": ("General Cargo", "Glass & Glassware"),
    "71": ("General Cargo", "Precious Metals & Jewellery"),
    "74": ("General Cargo", "Copper & Articles Thereof"),
    "76": ("General Cargo", "Aluminium & Articles Thereof")
}

def derive_category_subcategory(top_hs_str, trade_lane_str):
    if top_hs_str and str(top_hs_str).strip() != "":
        codes = [re.sub(r"[^0-9]", "", c) for c in str(top_hs_str).split(",")]
        for c in codes:
            if len(c) >= 2:
                ch = c[:2]
                if ch in CHAPTER_MAP:
                    return CHAPTER_MAP[ch]

    tl = str(trade_lane_str or "").upper()
    if "PER" in tl:
        return ("Perishables", "Perishables (Air Cargo)")
    elif "FABRIC" in tl:
        return ("Textiles & Apparel", "Fabrics & Technical Textiles")
    elif "YARN" in tl:
        return ("Textiles & Apparel", "Synthetic Filaments & Yarns")
    elif "EV" in tl:
        return ("Automotive & EV", "EV & 2-Wheeler Components")
    elif "GCR" in tl:
        return ("General Cargo", "General Cargo (Air)")

    return ("General Cargo", "General Cargo")

def enrich_valid_contacts_master():
    print("\n--- 1. ENRICHING VALID CONTACTS MASTER DIRECTORY ---")
    t0 = time.time()
    wb = openpyxl.load_workbook(VALID_MASTER_PATH)
    print(f"Loaded {VALID_MASTER_PATH} in {time.time()-t0:.2f}s")

    cat_counter = Counter()

    for sname in ["All Cleaned Contacts", "Exporters", "Importers"]:
        if sname not in wb.sheetnames:
            continue
        ws = wb[sname]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = list(rows[0])
        
        # Check if already enriched
        if "Category" in header and "Subcategory" in header:
            print(f"Sheet {sname} already has Category and Subcategory columns, skipping insertion.")
            continue

        c_idx = header.index("Company Name")
        r_idx = header.index("Role")
        tl_idx = header.index("Trade Lanes / Sectors")
        hs_idx = header.index("Top HS Codes")

        # Build new sheet content
        new_header = header[:2] + ["Category", "Subcategory"] + header[2:]
        
        enriched_rows = [new_header]
        for r in rows[1:]:
            cat, subcat = derive_category_subcategory(r[hs_idx], r[tl_idx])
            cat_counter[cat] += 1
            new_row = list(r[:2]) + [cat, subcat] + list(r[2:])
            enriched_rows.append(new_row)

        # Clear and rewrite sheet
        ws.delete_rows(1, ws.max_row)
        for r_data in enriched_rows:
            ws.append(r_data)

        # Header styling
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        for col_num in range(1, len(new_header) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font

        print(f"Enriched {sname}: {len(enriched_rows)-1} rows with Category & Subcategory.")

    # Update Audit & Metrics Summary
    if "Audit & Metrics Summary" in wb.sheetnames:
        summary_ws = wb["Audit & Metrics Summary"]
        summary_ws.append([])
        summary_ws.append(["LOGISTICS VERTICALS & SUBCATEGORIES ENRICHMENT", None])
        summary_ws.append(["Enrichment Execution Date", time.strftime("%Y-%m-%d %H:%M:%S")])
        summary_ws.append(["New Columns Added", "Category, Subcategory (Columns 3 & 4)"])
        for cat, cnt in cat_counter.most_common():
            summary_ws.append([f"Category: {cat}", f"{cnt} entity assignments"])

    t1 = time.time()
    wb.save(VALID_MASTER_PATH)
    print(f"Saved enriched Valid Contacts Master in {time.time()-t1:.2f}s")

def aggregate_and_enrich_missing_contacts():
    print("\n--- 2. AGGREGATING & ENRICHING MISSING CONTACTS MASTER DIRECTORY ---")
    t0 = time.time()
    missing_files = glob.glob(f"{DATAMYNE_DIR}/Missing Contacts/**/*.xlsx", recursive=True)
    missing_files = [f for f in missing_files if "Master" not in f]
    print(f"Found {len(missing_files)} transaction files in Missing Contacts.")

    # Data structure to aggregate entities
    # key: canonical upper company name
    entities = defaultdict(lambda: {
        "roles": set(),
        "tax_ids": Counter(),
        "cities": Counter(),
        "addresses": Counter(),
        "pin_codes": Counter(),
        "directors": Counter(),
        "trade_lanes": set(),
        "hs_codes": Counter(),
        "total_shipments": 0,
        "total_value_usd": 0.0,
        "latest_date": ""
    })

    total_records = 0
    for idx, fpath in enumerate(missing_files, 1):
        wb = openpyxl.load_workbook(fpath, read_only=True)
        ws = wb.active
        h = next(ws.iter_rows(values_only=True), None)
        if not h:
            continue
        
        # Derive trade lane corridor from file path
        lane = "GCR/Air"
        if "PER" in fpath:
            lane = "PER/AIR"
        elif "Fabric" in fpath:
            lane = "EX-China_INDIA/Fabric"
        elif "Yarn" in fpath:
            lane = "EX-China_INDIA/Yarn"
        elif "EV" in fpath:
            lane = "EX-China_INDIA/EV (2-Wheeler)"

        if "Importer" in h:
            role = "Importer"
            name_idx = h.index("Importer")
            tax_idx = h.index("Importer Tax ID") if "Importer Tax ID" in h else -1
            addr_idx = h.index("Importer Address") if "Importer Address" in h else -1
            city_idx = h.index("Importer City") if "Importer City" in h else -1
            zip_idx = h.index("Importer Zip Code") if "Importer Zip Code" in h else -1
            dir_idx = h.index("Importer Directors") if "Importer Directors" in h else -1
            hs_idx = h.index("HS Code") if "HS Code" in h else -1
            date_idx = h.index("Date") if "Date" in h else -1
            val_idx = h.index("Total Value (USD)") if "Total Value (USD)" in h else -1
        elif "Exporter" in h:
            role = "Exporter"
            name_idx = h.index("Exporter")
            tax_idx = h.index("Exporter Tax ID") if "Exporter Tax ID" in h else -1
            addr_idx = h.index("Exporter Address") if "Exporter Address" in h else -1
            city_idx = h.index("Exporter City") if "Exporter City" in h else -1
            zip_idx = h.index("Exporter Zip Code") if "Exporter Zip Code" in h else -1
            dir_idx = h.index("Exporter Directors") if "Exporter Directors" in h else -1
            hs_idx = h.index("HS Code") if "HS Code" in h else -1
            date_idx = h.index("Date") if "Date" in h else -1
            val_idx = h.index("FOB Value (USD)") if "FOB Value (USD)" in h else -1
        else:
            continue

        for r in ws.iter_rows(values_only=True):
            raw_name = r[name_idx]
            if not raw_name or not str(raw_name).strip():
                continue
            name = str(raw_name).strip().upper()
            total_records += 1

            ent = entities[name]
            ent["roles"].add(role)
            ent["trade_lanes"].add(lane)
            ent["total_shipments"] += 1

            if tax_idx >= 0 and r[tax_idx]:
                ent["tax_ids"][str(r[tax_idx]).strip()] += 1
            if city_idx >= 0 and r[city_idx]:
                ent["cities"][str(r[city_idx]).strip()] += 1
            if addr_idx >= 0 and r[addr_idx]:
                ent["addresses"][str(r[addr_idx]).strip()] += 1
            if zip_idx >= 0 and r[zip_idx]:
                ent["pin_codes"][str(r[zip_idx]).strip()] += 1
            if dir_idx >= 0 and r[dir_idx]:
                ent["directors"][str(r[dir_idx]).strip()] += 1
            if hs_idx >= 0 and r[hs_idx]:
                hscode = re.sub(r"[^0-9]", "", str(r[hs_idx]).strip())
                if hscode:
                    ent["hs_codes"][hscode] += 1
            if val_idx >= 0 and r[val_idx]:
                try:
                    val = float(r[val_idx])
                    if val > 0:
                        ent["total_value_usd"] += val
                except (ValueError, TypeError):
                    pass
            if date_idx >= 0 and r[date_idx]:
                d_str = str(r[date_idx]).split()[0]
                if d_str > ent["latest_date"]:
                    ent["latest_date"] = d_str

        if idx % 50 == 0 or idx == len(missing_files):
            print(f"Scanned {idx}/{len(missing_files)} files... ({len(entities)} unique entities, {total_records} records)")

    print(f"\nCompleted aggregation in {time.time()-t0:.2f}s.")
    print(f"Total Unique Missing Contact Entities: {len(entities)}")

    # Create new Workbook for Missing Contacts Master
    out_wb = openpyxl.Workbook()
    # Default sheet
    ws_all = out_wb.active
    ws_all.title = "All Cleaned Missing Contacts"
    ws_exp = out_wb.create_sheet(title="Exporters")
    ws_imp = out_wb.create_sheet(title="Importers")
    ws_summary = out_wb.create_sheet(title="Audit & Metrics Summary")

    header = [
        "Company Name",
        "Role",
        "Category",
        "Subcategory",
        "Contact Type",
        "Phone",
        "Email",
        "Tax ID (IEC/PAN)",
        "City",
        "Address",
        "PIN Code",
        "Directors / Contact",
        "Trade Lanes / Sectors",
        "Top HS Codes",
        "Total Shipments",
        "Total FOB/CIF (USD)",
        "Latest Activity Date"
    ]

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")

    for ws in [ws_all, ws_exp, ws_imp]:
        ws.append(header)
        for col_num in range(1, len(header) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font

    cat_dist = Counter()
    role_dist = Counter()
    total_usd = 0.0
    total_tx = 0

    sorted_names = sorted(entities.keys())
    for name in sorted_names:
        ent = entities[name]
        
        # Determine consolidated role
        if "Exporter" in ent["roles"] and "Importer" in ent["roles"]:
            role_str = "Both"
        elif "Exporter" in ent["roles"]:
            role_str = "Exporter"
        else:
            role_str = "Importer"
        role_dist[role_str] += 1

        # Most frequent attributes
        tax_id = ent["tax_ids"].most_common(1)[0][0] if ent["tax_ids"] else ""
        city = ent["cities"].most_common(1)[0][0] if ent["cities"] else ""
        address = ent["addresses"].most_common(1)[0][0] if ent["addresses"] else ""
        pin = ent["pin_codes"].most_common(1)[0][0] if ent["pin_codes"] else ""
        directors = ent["directors"].most_common(1)[0][0] if ent["directors"] else ""
        
        trade_lanes_str = ", ".join(sorted(ent["trade_lanes"]))
        top_hs_list = [c for c, _ in ent["hs_codes"].most_common(3)]
        top_hs_str = ", ".join(top_hs_list)

        cat, subcat = derive_category_subcategory(top_hs_str, trade_lanes_str)
        cat_dist[cat] += 1

        shipments = ent["total_shipments"]
        val_usd = round(ent["total_value_usd"], 2)
        total_usd += val_usd
        total_tx += shipments
        latest_date = ent["latest_date"] or "2026-02-28"

        row_data = [
            name,
            role_str,
            cat,
            subcat,
            "Missing Phone & Email",
            "",
            "",
            tax_id,
            city,
            address,
            pin,
            directors,
            trade_lanes_str,
            top_hs_str,
            shipments,
            val_usd,
            latest_date
        ]

        ws_all.append(row_data)
        if role_str in ["Exporter", "Both"]:
            ws_exp.append(row_data)
        if role_str in ["Importer", "Both"]:
            ws_imp.append(row_data)

    # Populate Audit & Metrics Summary
    ws_summary.append(["MISSING CONTACTS CONSOLIDATED MASTER DIRECTORY AUDIT", None])
    ws_summary.append(["Generated Date / Time", time.strftime("%Y-%m-%d %H:%M:%S")])
    ws_summary.append(["Total Source Transaction Files Scanned", len(missing_files)])
    ws_summary.append(["Total Shipping Transaction Records Processed", total_tx])
    ws_summary.append(["Total Unique Business Entities Identified", len(entities)])
    ws_summary.append(["Total Traded Value Represented (USD)", f"${total_usd:,.2f}"])
    ws_summary.append([])
    ws_summary.append(["Trading Role Distribution", None])
    ws_summary.append(["Exporters Only", role_dist["Exporter"]])
    ws_summary.append(["Importers Only", role_dist["Importer"]])
    ws_summary.append(["Both Exporter & Importer", role_dist["Both"]])
    ws_summary.append([])
    ws_summary.append(["Logistics Vertical Distribution", None])
    for cat, cnt in cat_dist.most_common():
        pct = (cnt / len(entities)) * 100
        ws_summary.append([cat, f"{cnt} entities ({pct:.1f}%)"])

    t_save = time.time()
    out_wb.save(MISSING_MASTER_PATH)
    print(f"Saved {MISSING_MASTER_PATH} in {time.time()-t_save:.2f}s")
    print(f"All Cleaned Missing Contacts rows: {len(entities)}")
    print(f"Exporters rows: {ws_exp.max_row - 1}")
    print(f"Importers rows: {ws_imp.max_row - 1}")

def main():
    start_time = time.time()
    print("=" * 70)
    print("STARTING MASTER DIRECTORIES ENRICHMENT & CONSOLIDATION")
    print("=" * 70)
    
    enrich_valid_contacts_master()
    aggregate_and_enrich_missing_contacts()

    print("\n" + "=" * 70)
    print(f"PROCESS COMPLETED IN {time.time()-start_time:.2f}s")
    print("=" * 70)

if __name__ == "__main__":
    main()

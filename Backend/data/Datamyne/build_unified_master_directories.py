import os
import re
import csv
import time
from collections import Counter, defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

CHAPTER_MAP = {
    # Perishables
    "02": ("Perishables", "Meat & Edible Meat Offal"),
    "03": ("Perishables", "Seafood & Marine Products"),
    "04": ("Perishables", "Dairy & Edible Animal Products"),
    "06": ("Perishables", "Floriculture & Cut Flowers"),
    "07": ("Perishables", "Fresh Vegetables & Produce"),
    "08": ("Perishables", "Fresh Fruits"),
    "09": ("Perishables", "Spices, Coffee & Tea"),
    "12": ("Perishables", "Agricultural Seeds & Sowing"),
    "13": ("Perishables", "Botanical Extracts & Resins"),
    "20": ("Perishables", "Processed Fruits & Vegetables"),
    "21": ("Perishables", "Processed Foods & Edible Preparations"),
    "22": ("Perishables", "Beverages & Liquids"),

    # Pharma & Healthcare
    "30": ("Pharma & Healthcare", "Pharmaceutical Formulations"),
    "90": ("Pharma & Healthcare", "Medical & Surgical Devices"),

    # Chemicals & Materials
    "28": ("Chemicals & Materials", "Inorganic Chemicals"),
    "29": ("Chemicals & Materials", "Organic Chemicals & APIs"),
    "38": ("Chemicals & Materials", "Specialty Chemicals & Auxiliaries"),

    # Textiles & Apparel
    "50": ("Textiles & Apparel", "Silk & Silk Textiles"),
    "51": ("Textiles & Apparel", "Wool & Animal Hair"),
    "52": ("Textiles & Apparel", "Cotton & Cotton Textiles"),
    "54": ("Textiles & Apparel", "Synthetic Filaments & Yarns"),
    "55": ("Textiles & Apparel", "Man-made Staple Fibres"),
    "56": ("Textiles & Apparel", "Wadding, Felt & Non-wovens"),
    "58": ("Textiles & Apparel", "Special Woven Fabrics"),
    "59": ("Textiles & Apparel", "Fabrics & Technical Textiles"),
    "60": ("Textiles & Apparel", "Knitted or Crocheted Fabrics"),
    "61": ("Textiles & Apparel", "Apparel & Ready-Made Garments (Knitted)"),
    "62": ("Textiles & Apparel", "Apparel & Ready-Made Garments (Woven)"),
    "63": ("Textiles & Apparel", "Made-up Textile Articles"),

    # Engineering & Machinery
    "84": ("Engineering & Machinery", "Industrial Machinery & Boilers"),
    "72": ("Engineering & Machinery", "Iron & Steel"),
    "73": ("Engineering & Machinery", "Iron, Steel & Hardware"),
    "82": ("Engineering & Machinery", "Tools & Implements"),

    # Electronics & Electrical
    "85": ("Electronics & Electrical", "Electrical Equipment & Switchgear"),

    # Automotive & EV
    "87": ("Automotive & EV", "Motor Vehicles, EV & Auto Parts"),

    # General Cargo
    "39": ("General Cargo", "Plastics & Polymers"),
    "40": ("General Cargo", "Rubber & Rubber Products"),
    "48": ("General Cargo", "Paper, Paperboard & Packaging"),
    "70": ("General Cargo", "Glass & Glassware"),
    "71": ("General Cargo", "Precious Metals & Jewellery"),
    "74": ("General Cargo", "Copper & Articles Thereof"),
    "76": ("General Cargo", "Aluminium & Articles Thereof"),
    "94": ("General Cargo", "Furniture, Bedding & Lighting")
}

def derive_category_subcategory(top_hs_str, trade_lane_str):
    if top_hs_str and str(top_hs_str).strip() not in ["", "None"]:
        codes = [re.sub(r"[^0-9]", "", c) for c in str(top_hs_str).split(",")]
        for c in codes:
            if len(c) >= 2:
                ch = c[:2]
                if ch in CHAPTER_MAP:
                    return CHAPTER_MAP[ch]

    tl = str(trade_lane_str or "").upper()
    if "PER" in tl:
        return ("Perishables", "Perishables (Air Cargo)")
    elif "FABRIC" in tl:
        return ("Textiles & Apparel", "Fabrics & Technical Textiles")
    elif "YARN" in tl:
        return ("Textiles & Apparel", "Synthetic Filaments & Yarns")
    elif "EV" in tl:
        return ("Automotive & EV", "EV & 2-Wheeler Components")
    elif "GCR" in tl:
        return ("General Cargo", "General Cargo (Air)")
    elif "OCEAN" in tl or "MUNDRA" in tl:
        return ("General Cargo", "General Cargo (Ocean)")

    return ("General Cargo", "General Cargo")

def norm_name(s):
    if not s:
        return ""
    s = str(s).upper().strip()
    s = re.sub(r"^(M/S\.?|MESSRS\.?|M/S\s+|MS\.)\s*", "", s)
    s = re.sub(r"&", " AND ", s)
    s = re.sub(r"\bPVT\.?\s*LTD\.?\b", "PRIVATE LIMITED", s)
    s = re.sub(r"\bLTD\.\b", "LIMITED", s)
    s = re.sub(r"[^A-Z0-9]", "", s)
    return s

def norm_phone(p):
    if not p:
        return ""
    digits = re.sub(r"\D", "", str(p))
    if len(digits) >= 10:
        return digits[-10:]
    return ""

def norm_email(e):
    if not e:
        return []
    e = str(e).lower().strip()
    if "@" in e and len(e) > 5:
        emails = [re.sub(r"\s+", "", x) for x in re.split(r"[,;]", e)]
        return [x for x in emails if "@" in x]
    return []

def clean_phone_display(p):
    if not p: return ""
    p_str = str(p).strip()
    p_str = re.sub(r"^(None|nan|\+?91[-\s]?)", "", p_str)
    return p_str

def parse_date(d):
    if not d: return "2023-01-01"
    d_str = str(d).strip()
    if re.match(r"^\d{4}-\d{2}-\d{2}", d_str):
        return d_str[:10]
    # Handle serial dates
    try:
        val = float(d_str)
        # Excel date serial
        import datetime
        dt = datetime.datetime(1899, 12, 30) + datetime.timedelta(days=val)
        return dt.strftime("%Y-%m-%d")
    except:
        pass
    return d_str[:10]

def load_active_exclusion():
    active_keys = set()
    with open("Backend/data/processed/active_customers_leads.csv", "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for r in reader:
            active_keys.add(norm_name(r["Full Name"]))
    return active_keys

def build_directories():
    t0 = time.time()
    print("=" * 70)
    print("BUILDING UNIFIED DATAMYNE MASTER DIRECTORIES WITH AIR & OCEAN")
    print("=" * 70)

    active_customers = load_active_exclusion()
    print(f"Loaded {len(active_customers)} active JFLL customer accounts to exclude.")

    # Storage for merged company entities
    # key: norm_name
    companies = {}

    def ingest_record(rec, source_dataset, default_mode):
        cname = str(rec.get("company_name", "")).strip()
        if not cname or cname == "None":
            return

        nn = norm_name(cname)
        if nn in active_customers:
            return  # Strictly exclude active customers

        raw_phone = clean_phone_display(rec.get("phone", ""))
        raw_email = str(rec.get("email", "")).strip()
        if raw_email in ["None", "nan"]: raw_email = ""
        
        # Phone numbers list
        phone_list = [p.strip() for p in re.split(r"[/,]", raw_phone) if p.strip()] if raw_phone else []
        # Emails list
        email_list = norm_email(raw_email)

        # Traded metrics
        try:
            shipments = int(float(rec.get("total_shipments", 1) or 1))
        except:
            shipments = 1
        try:
            val_usd = float(rec.get("total_value_usd", 0.0) or 0.0)
        except:
            val_usd = 0.0

        activity_date = parse_date(rec.get("latest_activity_date", "2023-01-01"))

        # HS Codes
        hs_raw = str(rec.get("top_hs_codes", "")).strip()
        hs_codes = [h.strip() for h in hs_raw.split(",") if h.strip() and h.strip() != "None"] if hs_raw else []

        # Trade Lanes
        tl_raw = str(rec.get("trade_lanes", "")).strip()
        trade_lanes = set([t.strip() for t in tl_raw.split(",") if t.strip() and t.strip() != "None"]) if tl_raw else set()
        if not trade_lanes and default_mode:
            trade_lanes.add(f"{default_mode}/Export")

        # Mode determination
        modes = set()
        if default_mode == "Air":
            modes.add("Air")
        elif default_mode == "Ocean":
            modes.add("Ocean")
        else:
            # Check trade lanes
            for tl in trade_lanes:
                tl_upper = tl.upper()
                if "AIR" in tl_upper or "GCR" in tl_upper or "PER" in tl_upper:
                    modes.add("Air")
                if "OCEAN" in tl_upper or "SEA" in tl_upper or "CHINA" in tl_upper or "MUNDRA" in tl_upper:
                    modes.add("Ocean")
        if not modes:
            modes.add("Air")

        if nn not in companies:
            companies[nn] = {
                "company_name": cname,
                "role": rec.get("role", "Exporter"),
                "modes": set(modes),
                "category": rec.get("category", ""),
                "subcategory": rec.get("subcategory", ""),
                "phones": list(phone_list),
                "emails": list(email_list),
                "tax_id": str(rec.get("tax_id", "")).strip() if rec.get("tax_id") and str(rec.get("tax_id")).strip() != "None" else "",
                "city": str(rec.get("city", "")).strip() if rec.get("city") and str(rec.get("city")).strip() != "None" else "",
                "address": str(rec.get("address", "")).strip() if rec.get("address") and str(rec.get("address")).strip() != "None" else "",
                "pin_code": str(rec.get("pin_code", "")).strip() if rec.get("pin_code") and str(rec.get("pin_code")).strip() != "None" else "",
                "directors": str(rec.get("directors", "")).strip() if rec.get("directors") and str(rec.get("directors")).strip() != "None" else "",
                "trade_lanes": set(trade_lanes),
                "hs_codes": list(hs_codes),
                "total_shipments": shipments,
                "total_value_usd": val_usd,
                "latest_activity_date": activity_date,
                "sources": {source_dataset}
            }
        else:
            # Merge into existing company
            c = companies[nn]
            c["modes"].update(modes)
            c["sources"].add(source_dataset)
            c["total_shipments"] += shipments
            c["total_value_usd"] += val_usd
            if activity_date > c["latest_activity_date"]:
                c["latest_activity_date"] = activity_date
            
            # Combine phones & emails
            for p in phone_list:
                if p not in c["phones"]: c["phones"].append(p)
            for em in email_list:
                if em not in c["emails"]: c["emails"].append(em)

            # Combine trade lanes & HS codes
            c["trade_lanes"].update(trade_lanes)
            for h in hs_codes:
                if h not in c["hs_codes"]: c["hs_codes"].append(h)

            # Fill missing attributes
            if not c["tax_id"] and rec.get("tax_id") and str(rec.get("tax_id")).strip() != "None":
                c["tax_id"] = str(rec.get("tax_id")).strip()
            if not c["city"] and rec.get("city") and str(rec.get("city")).strip() != "None":
                c["city"] = str(rec.get("city")).strip()
            if not c["address"] and rec.get("address") and str(rec.get("address")).strip() != "None":
                c["address"] = str(rec.get("address")).strip()
            if not c["pin_code"] and rec.get("pin_code") and str(rec.get("pin_code")).strip() != "None":
                c["pin_code"] = str(rec.get("pin_code")).strip()
            if not c["directors"] and rec.get("directors") and str(rec.get("directors")).strip() != "None":
                c["directors"] = str(rec.get("directors")).strip()

    # 1. Ingest Original Valid Contacts
    print("Ingesting Original Valid Contacts Master Directory...")
    wb = openpyxl.load_workbook("Backend/data/Datamyne/Valid Contacts/Master_Cleaned_Contacts_Directory.xlsx", read_only=True)
    ws = wb["All Cleaned Contacts"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c) for c in rows[0]]
    for r in rows[1:]:
        d = dict(zip(header, r))
        ingest_record({
            "company_name": d.get("Company Name"),
            "role": d.get("Role"),
            "category": d.get("Category"),
            "subcategory": d.get("Subcategory"),
            "phone": d.get("Phone"),
            "email": d.get("Email"),
            "tax_id": d.get("Tax ID (IEC/PAN)"),
            "city": d.get("City"),
            "address": d.get("Address"),
            "pin_code": d.get("PIN Code"),
            "directors": d.get("Directors / Contact"),
            "trade_lanes": d.get("Trade Lanes / Sectors"),
            "top_hs_codes": d.get("Top HS Codes"),
            "total_shipments": d.get("Total Shipments"),
            "total_value_usd": d.get("Total FOB/CIF (USD)"),
            "latest_activity_date": d.get("Latest Activity Date")
        }, "DM1_Valid", None)

    # 2. Ingest Original Missing Contacts
    print("Ingesting Original Missing Contacts Master Directory...")
    wb = openpyxl.load_workbook("Backend/data/Datamyne/Missing Contacts/Master_Cleaned_Missing_Contacts_Directory.xlsx", read_only=True)
    ws = wb["All Cleaned Missing Contacts"]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c) for c in rows[0]]
    for r in rows[1:]:
        d = dict(zip(header, r))
        ingest_record({
            "company_name": d.get("Company Name"),
            "role": d.get("Role"),
            "category": d.get("Category"),
            "subcategory": d.get("Subcategory"),
            "phone": d.get("Phone"),
            "email": d.get("Email"),
            "tax_id": d.get("Tax ID (IEC/PAN)"),
            "city": d.get("City"),
            "address": d.get("Address"),
            "pin_code": d.get("PIN Code"),
            "directors": d.get("Directors / Contact"),
            "trade_lanes": d.get("Trade Lanes / Sectors"),
            "top_hs_codes": d.get("Top HS Codes"),
            "total_shipments": d.get("Total Shipments"),
            "total_value_usd": d.get("Total FOB/CIF (USD)"),
            "latest_activity_date": d.get("Latest Activity Date")
        }, "DM1_Missing", None)

    # 3. Ingest DataMyne_2 Air Master (from backup)
    print("Ingesting DataMyne_2 Air Master Directory...")
    wb = openpyxl.load_workbook("Backend/data/.backups/2026-10-17-pre-merge/Datamyne_2/Air/Master_Cleaned_Air_Contacts.xlsx", read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c) for c in rows[0]]
    for r in rows[1:]:
        d = dict(zip(header, r))
        ingest_record({
            "company_name": d.get("Company Name"),
            "role": d.get("Role", "Exporter"),
            "phone": d.get("Phone"),
            "email": d.get("Email"),
            "tax_id": d.get("Tax ID (IEC/PAN)"),
            "city": d.get("City"),
            "address": d.get("Address"),
            "pin_code": d.get("PIN Code"),
            "directors": d.get("Directors / Contact"),
            "trade_lanes": "Air/Export",
            "top_hs_codes": d.get("Top HS Codes"),
            "total_shipments": d.get("File Appearances", 1),
            "total_value_usd": d.get("Total FOB/CIF (USD)"),
            "latest_activity_date": d.get("Latest Activity Date")
        }, "DM2_Air", "Air")

    # 4. Ingest DataMyne_2 Ocean Master (from backup)
    print("Ingesting DataMyne_2 Ocean Master Directory...")
    wb = openpyxl.load_workbook("Backend/data/.backups/2026-10-17-pre-merge/Datamyne_2/Ocean/Master_Cleaned_Ocean_Contacts.xlsx", read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(c) for c in rows[0]]
    for r in rows[1:]:
        d = dict(zip(header, r))
        ingest_record({
            "company_name": d.get("Company Name"),
            "role": d.get("Role", "Exporter"),
            "phone": d.get("Phone"),
            "email": d.get("Email"),
            "tax_id": d.get("Tax ID (IEC/PAN)"),
            "city": d.get("City"),
            "address": d.get("Address"),
            "pin_code": d.get("PIN Code"),
            "directors": d.get("Directors / Contact"),
            "trade_lanes": "Ocean/Export",
            "top_hs_codes": d.get("Top HS Codes"),
            "total_shipments": d.get("File Appearances", 1),
            "total_value_usd": d.get("Total FOB/CIF (USD)"),
            "latest_activity_date": d.get("Latest Activity Date")
        }, "DM2_Ocean", "Ocean")

    print(f"\nTotal unique companies after deduplication & merging: {len(companies):,}")

    # Final record preparation
    records = []
    mode_counts = Counter()
    dataset_counts = Counter()

    for nn, c in companies.items():
        # Transport mode string
        if "Air" in c["modes"] and "Ocean" in c["modes"]:
            transport_mode = "Air & Ocean"
        elif "Ocean" in c["modes"]:
            transport_mode = "Ocean"
        else:
            transport_mode = "Air"
        mode_counts[transport_mode] += 1

        # Phones and Emails formatted
        phone_str = " / ".join(c["phones"]) if c["phones"] else ""
        email_str = ", ".join(c["emails"]) if c["emails"] else ""

        # Contact Type
        has_p = bool(phone_str)
        has_e = bool(email_str)
        if has_p and has_e:
            contact_type = "Both Phone & Email"
            dataset_type = "valid"
        elif has_p:
            contact_type = "Phone Only"
            dataset_type = "valid"
        elif has_e:
            contact_type = "Email Only"
            dataset_type = "valid"
        else:
            contact_type = "Missing Phone & Email"
            dataset_type = "missing"
        dataset_counts[dataset_type] += 1

        # Derive category if missing
        cat = c["category"]
        subcat = c["subcategory"]
        top_hs_str = ", ".join(c["hs_codes"][:4]) if c["hs_codes"] else ""
        tl_str = ", ".join(sorted(c["trade_lanes"])) if c["trade_lanes"] else ""
        if not cat or not subcat or cat == "General Cargo":
            derived_cat, derived_subcat = derive_category_subcategory(top_hs_str, tl_str)
            if not cat: cat = derived_cat
            if not subcat: subcat = derived_subcat

        rec_out = {
            "Company Name": c["company_name"],
            "Role": c["role"],
            "Transport Mode": transport_mode,
            "Category": cat,
            "Subcategory": subcat,
            "Contact Type": contact_type,
            "Phone": phone_str,
            "Email": email_str,
            "Tax ID (IEC/PAN)": c["tax_id"],
            "City": c["city"],
            "Address": c["address"],
            "PIN Code": c["pin_code"],
            "Directors / Contact": c["directors"],
            "Trade Lanes / Sectors": tl_str,
            "Top HS Codes": top_hs_str,
            "Total Shipments": c["total_shipments"],
            "Total FOB/CIF (USD)": round(c["total_value_usd"], 2),
            "Latest Activity Date": c["latest_activity_date"],
            "dataset_type": dataset_type,
            "modes_set": c["modes"]
        }
        records.append(rec_out)

    print("\nTransport Mode Breakdown:")
    for m, cnt in mode_counts.items():
        print(f"  • {m}: {cnt:,} leads")

    print("\nContact Readiness Breakdown:")
    for ds, cnt in dataset_counts.items():
        print(f"  • {ds.capitalize()}: {cnt:,} leads")

    # Function to save styled Excel workbook
    def save_styled_excel(file_path, sheet_name, data_list):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name

        headers = [
            "Company Name",
            "Role",
            "Transport Mode",
            "Category",
            "Subcategory",
            "Contact Type",
            "Phone",
            "Email",
            "Tax ID (IEC/PAN)",
            "City",
            "Address",
            "PIN Code",
            "Directors / Contact",
            "Trade Lanes / Sectors",
            "Top HS Codes",
            "Total Shipments",
            "Total FOB/CIF (USD)",
            "Latest Activity Date"
        ]

        ws.append(headers)

        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        data_font = Font(name="Arial", size=9)

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.row_dimensions[1].height = 24

        for row_idx, r in enumerate(data_list, start=2):
            row_vals = [r.get(h, "") for h in headers]
            ws.append(row_vals)
            # Style row
            ws.cell(row=row_idx, column=16).number_format = "#,##0"
            ws.cell(row=row_idx, column=17).number_format = "$#,##0.00"

        # Autofit column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col[:100]:  # sample first 100 rows for speed
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(12, min(max_len + 3, 40))

        os.makedirs(os.path.dirname(file_path), exist_ok=True)
        wb.save(file_path)
        print(f"  ✓ Saved {sheet_name} ({len(data_list):,} rows) -> {file_path}")

    print("\n--- Generating Master Directories ---")
    # 1. Air Master Directory (Air & Air & Ocean)
    air_records = [r for r in records if "Air" in r["modes_set"]]
    save_styled_excel(
        "Backend/data/Datamyne/Air/Master_Cleaned_Air_Contacts.xlsx",
        "Air Freight Logistics Leads",
        air_records
    )

    # 2. Ocean Master Directory (Ocean & Air & Ocean)
    ocean_records = [r for r in records if "Ocean" in r["modes_set"]]
    save_styled_excel(
        "Backend/data/Datamyne/Ocean/Master_Cleaned_Ocean_Contacts.xlsx",
        "Ocean Freight Logistics Leads",
        ocean_records
    )

    # 3. Valid Contacts Directory (With Contacts, all modes)
    valid_records = [r for r in records if r["dataset_type"] == "valid"]
    save_styled_excel(
        "Backend/data/Datamyne/Valid Contacts/Master_Cleaned_Contacts_Directory.xlsx",
        "All Cleaned Contacts",
        valid_records
    )

    # 4. Missing Contacts Directory (Without Contacts, all modes)
    missing_records = [r for r in records if r["dataset_type"] == "missing"]
    save_styled_excel(
        "Backend/data/Datamyne/Missing Contacts/Master_Cleaned_Missing_Contacts_Directory.xlsx",
        "All Cleaned Missing Contacts",
        missing_records
    )

    print(f"\nAll Master Directories generated in {time.time()-t0:.2f}s!")
    print("=" * 70)

if __name__ == "__main__":
    build_directories()

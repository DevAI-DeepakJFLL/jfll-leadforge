import os
import re
import csv
import time
from itertools import groupby
from operator import itemgetter
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

ACTIVE_CUSTOMERS_PATH = "Backend/data/processed/active_customers_leads.csv"
DM2_DIR = "Backend/data/Datamyne_2"
AUDIT_CSV_PATH = "Backend/data/Datamyne/DUPLICATES_REMOVED_DATAMYNE2_AUDIT.csv"

def norm_name(s):
    if not s:
        return ""
    s = str(s).upper().strip()
    s = re.sub(r"^(M/S\.?|MESSRS\.?|M/S\s+|MS\.)\s*", "", s)
    s = re.sub(r"&", " AND ", s)
    s = re.sub(r"\bPVT\.?\s*LTD\.?\b", "PRIVATE LIMITED", s)
    s = re.sub(r"\bLTD\.\b", "LIMITED", s)
    s = re.sub(r"[^A-Z0-9]", "", s)
    return s

def norm_phone(p):
    if not p:
        return ""
    digits = re.sub(r"\D", "", str(p))
    if len(digits) >= 10:
        return digits[-10:]
    return ""

def norm_email(e):
    if not e:
        return []
    e = str(e).lower().strip()
    if "@" in e and len(e) > 5:
        emails = [re.sub(r"\s+", "", x) for x in re.split(r"[,;]", e)]
        return [x for x in emails if "@" in x]
    return []

def load_active_customers():
    leads = []
    with open(ACTIVE_CUSTOMERS_PATH, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            name = row["Full Name"].strip()
            leads.append({
                "name": name,
                "n_name": norm_name(name),
                "phone": norm_phone(row.get("Mobile Number", "")),
                "emails": norm_email(row.get("Email", "")),
                "raw_phone": row.get("Mobile Number", ""),
                "raw_email": row.get("Email", "")
            })
    return leads

def match_lead(c_name, c_phone, c_emails, active_leads):
    n_cn = norm_name(c_name)
    n_ph = norm_phone(c_phone)
    matched_leads = []
    
    for lead in active_leads:
        reasons = []
        if n_cn and n_cn == lead["n_name"]:
            reasons.append("Name Match")
        if n_ph and lead["phone"] and n_ph == lead["phone"]:
            reasons.append("Phone Match")
        for em in c_emails:
            if em in lead["emails"]:
                reasons.append("Email Match")
                break
        if reasons:
            matched_leads.append((lead, " + ".join(reasons)))
            
    return matched_leads

def delete_row_chunks(ws, rows_to_delete):
    rows_to_delete = sorted(set(rows_to_delete), reverse=True)
    chunks = []
    for k, g in groupby(enumerate(rows_to_delete), lambda ix: ix[0] + ix[1]):
        chunk = [x[1] for x in g]
        chunks.append((chunk[-1], len(chunk)))
    for start_row, count in chunks:
        ws.delete_rows(start_row, count)

def process_file(file_path, active_leads):
    try:
        wb = openpyxl.load_workbook(file_path)
    except Exception as e:
        return file_path, 0, [], f"Error loading: {e}"

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return file_path, 0, [], "Empty file"

    header = [str(c) for c in rows[0]]
    if "Company Name" in header:
        name_idx = header.index("Company Name")
        phone_idx = header.index("Phone") if "Phone" in header else -1
        email_idx = header.index("Email") if "Email" in header else -1
        category = "Master Directory"
    elif "Importer" in header:
        name_idx = header.index("Importer")
        phone_idx = header.index("Importer Phone") if "Importer Phone" in header else -1
        email_idx = next((i for i, h in enumerate(header) if "importer e-mail" in h.lower() or "importer email" in h.lower()), -1)
        category = "Raw Contacts (Importer)"
    elif "Exporter" in header:
        name_idx = header.index("Exporter")
        phone_idx = header.index("Exporter Phone") if "Exporter Phone" in header else -1
        email_idx = next((i for i, h in enumerate(header) if "exporter e-mail" in h.lower() or "exporter email" in h.lower()), -1)
        category = "Raw Contacts (Exporter)"
    else:
        return file_path, 0, [], "Unknown schema"

    to_delete = []
    audit_entries = []

    for row_idx, r in enumerate(rows[1:], start=2):
        c_name = r[name_idx] if name_idx >= 0 and len(r) > name_idx else None
        c_phone = r[phone_idx] if phone_idx >= 0 and len(r) > phone_idx else None
        c_emails = norm_email(r[email_idx]) if (email_idx >= 0 and len(r) > email_idx) else []

        matches = match_lead(c_name, c_phone, c_emails, active_leads)
        if matches:
            to_delete.append(row_idx)
            for lead, reason in matches:
                audit_entries.append({
                    "Category": category,
                    "File Path": file_path,
                    "Original Row Number": row_idx,
                    "DataMyne Entity Name": c_name or "",
                    "Matched JFLL Customer Name": lead["name"],
                    "Match Reason": reason,
                    "DataMyne Phone": str(c_phone) if c_phone else "",
                    "JFLL Phone": lead["raw_phone"],
                    "DataMyne Email": str(r[email_idx]) if (email_idx >= 0 and len(r) > email_idx and r[email_idx]) else "",
                    "JFLL Email": lead["raw_email"]
                })

    if to_delete:
        delete_row_chunks(ws, to_delete)
        wb.save(file_path)

    return file_path, len(to_delete), audit_entries, None

def run_purge():
    t0 = time.time()
    print("=" * 70)
    print("PURGING ACTIVE JFLL CUSTOMERS FROM DATAMYNE_2")
    print("=" * 70)

    active_leads = load_active_customers()
    print(f"Loaded {len(active_leads)} active JFLL customer accounts.")

    all_files = []
    for root, dirs, files in os.walk(DM2_DIR):
        for f in files:
            if f.endswith(".xlsx"):
                all_files.append(os.path.join(root, f))

    print(f"Found {len(all_files)} total Excel files in {DM2_DIR}.")

    total_deleted = 0
    all_audit_entries = []

    for i, fp in enumerate(all_files, start=1):
        fp_res, del_count, audit_entries, err = process_file(fp, active_leads)
        if err:
            print(f"[{i}/{len(all_files)}] ERROR in {fp}: {err}")
        elif del_count > 0:
            print(f"[{i}/{len(all_files)}] {os.path.basename(fp)}: Deleted {del_count} rows.")
            total_deleted += del_count
            all_audit_entries.extend(audit_entries)

    print(f"\nTotal rows purged across all DataMyne_2 files: {total_deleted}")
    print(f"Total audit entries generated: {len(all_audit_entries)}")

    # Write audit CSV
    fieldnames = [
        "Category",
        "File Path",
        "Original Row Number",
        "DataMyne Entity Name",
        "Matched JFLL Customer Name",
        "Match Reason",
        "DataMyne Phone",
        "JFLL Phone",
        "DataMyne Email",
        "JFLL Email"
    ]

    with open(AUDIT_CSV_PATH, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for entry in all_audit_entries:
            writer.writerow(entry)

    print(f"Audit log successfully written to: {AUDIT_CSV_PATH}")
    print(f"Purge completed in {time.time()-t0:.2f}s")
    print("=" * 70)

if __name__ == "__main__":
    run_purge()
